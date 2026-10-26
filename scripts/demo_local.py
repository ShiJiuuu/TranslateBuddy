#!/usr/bin/env python3
"""不连 Dify 也能跑通的端到端演示：生成样例 → 解析 → 假翻译 → 就地回写。

真跑时把 stub_invoke 换成 session.model.llm.invoke 即可（见 tools/translate_document.py）。

用法:
    python scripts/demo_local.py                 # docx
    python scripts/demo_local.py --format xlsx   # xlsx
    python scripts/demo_local.py --src my.docx --out out.docx
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import RGBColor

from backend.formats import open_document
from backend.qa import check
from backend.tags import strip_tags
from backend.translator import translate_segments

DICT = {
    "Quarterly Report": "季度报告",
    "Revenue grew by 18 percent": "营收同比增长 18%",
    "Operating expenses remained flat": "运营费用基本持平",
    "enterprise segment": "企业客户",
    "Confidential": "机密",
    "Metric": "指标",
    "Value": "数值",
    "Note": "备注",
    "Revenue": "营收",
    "Cost": "成本",
    "Margin": "利润率",
    "Enterprise demand": "企业需求",
}


def stub_invoke(system: str, user: str) -> str:
    payload = json.loads(user)
    items = []
    for item in payload["items"]:
        text = item["text"]
        for src in sorted(DICT, key=len, reverse=True):
            text = text.replace(src, DICT[src])
        items.append({"id": item["id"], "text": text})
    return json.dumps({"items": items}, ensure_ascii=False)


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Quarterly Report 2024", level=1)
    p = doc.add_paragraph()
    run = p.add_run("Revenue grew by 18 percent")
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p.add_run(" compared with the same period last year, driven by strong demand in the ")
    italic = p.add_run("enterprise segment")
    italic.italic = True
    p.add_run(".")
    doc.add_paragraph("Operating expenses remained flat.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "1,240"
    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(path)


def build_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Quarterly Report 2024"
    ws["A1"].font = Font(size=16, bold=True)
    for col, text in enumerate(["Metric", "Value", "Note"], start=1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    ws["A4"], ws["B4"], ws["C4"] = "Revenue", 1240, "Enterprise demand"
    ws["A5"], ws["B5"] = "Cost", 880
    ws["A6"], ws["B6"] = "Margin", "=B4-B5"
    ws["B4"].number_format = '#,##0 "USD"'
    wb.save(path)


BUILDERS = {"docx": build_docx, "xlsx": build_xlsx}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--src", type=Path)
    parser.add_argument("--format", choices=sorted(BUILDERS), default="docx")
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    fmt = (args.src.suffix.lstrip(".") if args.src else args.format)
    out = args.out or Path(f"demo_out.{fmt}")

    with tempfile.TemporaryDirectory() as td:
        src = args.src
        if src is None:
            src = Path(td) / f"sample.{fmt}"
            BUILDERS[fmt](src)
            print("已生成样例：", src)

        doc = open_document(src)
        print(f"解析：{len(doc.segments)} 段可翻译文本")

        translations = translate_segments(doc.segments, stub_invoke, "zh-Hans")
        fallback = 0
        issues = []
        for seg in doc.segments:
            text = translations.get(seg.id)
            if not text:
                continue
            if not doc.apply(seg.id, text):
                fallback += 1
            issues.extend(check(seg.plain, strip_tags(text)))

        doc.save(out)
        print(f"回写完成：{len(translations)}/{len(doc.segments)} 段，格式降级 {fallback} 次，QA {len(issues)} 条")
        print("输出：", out.resolve())

    if fmt == "docx":
        result = Document(out)
        for paragraph in result.paragraphs[:3]:
            if not paragraph.text.strip():
                continue
            print("\n段落：", paragraph.text)
            for run in paragraph.runs:
                if run.text:
                    color = run.font.color.rgb if run.font.color and run.font.color.rgb else None
                    print(f"   run={run.text!r} bold={run.bold} italic={run.italic} color={color}")
    else:
        from openpyxl import load_workbook

        ws = load_workbook(out).active
        for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
            print("  ", row)


if __name__ == "__main__":
    main()
