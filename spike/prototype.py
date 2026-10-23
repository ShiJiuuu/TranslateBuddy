"""Feasibility spike: format-preserving in-place translation for docx / xlsx / pdf.

Translation is stubbed with a tiny dictionary so the spike isolates the *format* problem,
which is the part the plugin is supposed to own. A real run swaps `tr()` for an LLM call.
"""
from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

import pymupdf
from lxml import etree
from openpyxl import load_workbook

HERE = Path(__file__).parent
SAMPLES = HERE / "samples"
OUT = HERE / "out"
OUT.mkdir(exist_ok=True)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}

FULL = {
    "Quarterly Report 2024": "2024 年季度报告",
    "Revenue grew by 18 percent compared with the same period last year, driven by strong demand in the enterprise segment.": "受企业客户需求强劲推动，营收较去年同期增长 18%。",
    "Operating expenses remained flat as the team continued to invest in automation and customer support.": "团队持续投入自动化与客户支持，运营费用基本持平。",
}
PHRASES = {
    "Revenue grew by 18 percent": "营收同比增长 18%",
    " compared with the same period last year, driven by strong demand in the ": "，这一增长主要来自",
    "enterprise segment": "企业客户",
    "Operating expenses remained flat": "运营费用基本持平",
    " as the team continued to invest in ": "，团队继续投入",
    "automation and customer support": "自动化与客户支持",
    "Quarterly Report 2024": "2024 年季度报告",
    "Metric": "指标",
    "Value": "数值",
    "Revenue": "营收",
    "Cost": "成本",
    "Margin": "利润率",
    "Confidential": "机密",
    "Page 1": "第 1 页",
    "Outlook": "展望",
    "We expect growth to continue in the next two quarters.": "我们预计未来两个季度将继续增长。",
}


def tr(text: str) -> str:
    """Stub translator: sentence dict first, then longest-phrase replacement."""
    s = text.strip()
    if not s:
        return text
    if s in FULL:
        out = FULL[s]
    elif s in PHRASES:
        out = PHRASES[s]
    else:
        out = text
        for k in sorted(PHRASES, key=len, reverse=True):
            out = out.replace(k, PHRASES[k])
    if out == text:
        return text
    lead = text[: len(text) - len(text.lstrip())]
    tail = text[len(text.rstrip()):]
    return lead + out.strip() + tail


# --------------------------------------------------------------------------- docx
def w(tag: str) -> str:
    return f"{{{W}}}{tag}"


def run_is_plain(r) -> bool:
    for child in r:
        if etree.QName(child).localname in ("rPr", "t"):
            continue
        return False
    return True


def run_sig(r) -> bytes:
    rpr = r.find(w("rPr"))
    return b"" if rpr is None else etree.tostring(rpr)


def run_text(r) -> str:
    return "".join(t.text or "" for t in r.findall(w("t")))


def set_run_text(r, text: str) -> None:
    ts = r.findall(w("t"))
    if not ts:
        t = etree.SubElement(r, w("t"))
        ts = [t]
    ts[0].text = text
    if text != text.strip():
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for extra in ts[1:]:
        r.remove(extra)


def translate_part(xml: bytes) -> tuple[bytes, int]:
    root = etree.fromstring(xml)
    touched = 0
    for p in root.iter(w("p")):
        # group consecutive plain runs that share the same run properties
        groups, cur = [], []
        for r in p.iter(w("r")):
            if not run_is_plain(r):
                if cur:
                    groups.append(cur)
                cur = []
                continue
            if cur and run_sig(cur[-1]) != run_sig(r):
                groups.append(cur)
                cur = []
            cur.append(r)
        if cur:
            groups.append(cur)

        for g in groups:
            src = "".join(run_text(r) for r in g)
            if not re.search(r"[A-Za-z]", src):
                continue
            dst = tr(src)
            if dst == src:
                continue
            set_run_text(g[0], dst)
            for r in g[1:]:
                set_run_text(r, "")
            touched += 1
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), touched


def docx_inplace(src: Path, dst: Path) -> dict:
    stats = {"parts": [], "segments": 0}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$", item.filename):
                data, n = translate_part(data)
                stats["segments"] += n
                stats["parts"].append(f"{item.filename}:{n}")
            zout.writestr(item, data)
    return stats


# --------------------------------------------------------------------------- xlsx
def xlsx_inplace(src: Path, dst: Path) -> dict:
    wb = load_workbook(src, rich_text=False)
    stats = {"cells": 0, "sheets": []}
    for ws in wb.worksheets:
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if c.data_type != "s":        # str cells only: skip formulas, numbers, dates
                    continue
                if not isinstance(c.value, str) or not re.search(r"[A-Za-z]", c.value):
                    continue
                new = tr(c.value)
                if new != c.value:
                    c.value = new
                    n += 1
        stats["cells"] += n
        stats["sheets"].append(f"{ws.title}:{n}")
    wb.save(dst)
    return stats


# --------------------------------------------------------------------------- pdf
def span_hex(span) -> str:
    return f"#{span['color'] & 0xFFFFFF:06x}"


def cluster_spans(spans):
    """Split a block into column cells when its lines have wide horizontal gaps."""
    spans = sorted(spans, key=lambda s: s["bbox"][0])
    clusters = []
    for s in spans:
        if clusters and s["bbox"][0] - clusters[-1][-1]["bbox"][2] < 24:
            clusters[-1].append(s)
        else:
            clusters.append([s])
    return clusters


def pdf_inplace(src: Path, dst: Path) -> dict:
    doc = pymupdf.open(src)
    stats = {"blocks": 0, "cells": 0, "pages": len(doc)}
    for page in doc:
        jobs = []
        for block in page.get_text("dict")["blocks"]:
            if block["type"] != 0:
                continue
            lines = block["lines"]
            flat = [sp for line in lines for sp in line["spans"] if sp["text"].strip()]
            if not flat:
                continue
            multi_col = any(len([s for s in line["spans"] if s["text"].strip()]) > 1 for line in lines)
            if multi_col:
                groups = cluster_spans(flat)
            else:
                groups = [flat]
            for g in groups:
                text = " ".join(s["text"].strip() for s in g)
                if not re.search(r"[A-Za-z]", text):
                    continue
                new = tr(text)
                if new == text:
                    continue
                rect = pymupdf.Rect(g[0]["bbox"])
                for s in g[1:]:
                    rect |= pymupdf.Rect(s["bbox"])
                jobs.append((rect, new, max(s["size"] for s in g), span_hex(g[0])))

        for rect, _, _, _ in jobs:
            page.add_redact_annot(rect, fill=False, cross_out=False)
        if jobs:
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                                  graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                                  text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        for rect, text, size, color in jobs:
            page.insert_htmlbox(
                rect + (-1, -3, 1, 3),
                f"<p>{text}</p>",
                css=f"* {{font-family: sans-serif; font-size: {size}px; color: {color}; margin: 0; line-height: 1.15;}}",
                scale_low=0.45,
            )
            stats["cells"] += 1
        stats["blocks"] += len(jobs)
    doc.subset_fonts()          # 不做这步：一页 PDF 会从 3KB 涨到 12MB
    doc.save(dst, garbage=4, deflate=True)
    return stats


if __name__ == "__main__":
    for name in ("sample.docx", "sample.xlsx", "sample.pdf"):
        shutil.copy(SAMPLES / name, OUT / f"orig_{name}")
    print("docx", docx_inplace(SAMPLES / "sample.docx", OUT / "zh_sample.docx"))
    print("xlsx", xlsx_inplace(SAMPLES / "sample.xlsx", OUT / "zh_sample.xlsx"))
    print("pdf ", pdf_inplace(SAMPLES / "sample.pdf", OUT / "zh_sample.pdf"))
