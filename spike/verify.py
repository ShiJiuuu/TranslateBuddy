import json
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

OUT = Path(__file__).parent / "out"
report = {}

# ---- docx
d = Document(OUT / "zh_sample.docx")
paras = []
for p in d.paragraphs:
    if p.text.strip():
        runs = [{"t": r.text, "b": bool(r.bold), "i": bool(r.italic),
                 "color": str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None}
                for r in p.runs if r.text]
        paras.append({"style": p.style.name, "runs": runs})
table = [[c.text for c in row.cells] for row in d.tables[0].rows]
report["docx"] = {
    "paragraphs": paras,
    "table": table,
    "header": d.sections[0].header.paragraphs[0].text,
    "footer": d.sections[0].footer.paragraphs[0].text,
}

# ---- xlsx
wb = load_workbook(OUT / "zh_sample.xlsx")
ws = wb["Report"]
report["xlsx"] = {
    "merged": [str(r) for r in ws.merged_cells.ranges],
    "a1": [ws["A1"].value, ws["A1"].font.size, ws["A1"].font.bold, ws["A1"].font.color.rgb if ws["A1"].font.color else None],
    "header_fill": [ws["A3"].value, ws["A3"].fill.fgColor.rgb, ws["A3"].font.bold],
    "cells": [[ws.cell(r, c).value for c in range(1, 4)] for r in range(4, 7)],
    "numfmt": [ws["B4"].number_format, ws["B6"].number_format],
    "formula": ws["B6"].value,
    "wrap": ws["C4"].alignment.wrap_text,
}

# ---- pdf
doc = pymupdf.open(OUT / "zh_sample.pdf")
report["pdf"] = {"text": doc[0].get_text().strip()[:400], "pages": len(doc)}
for name in ("orig_sample.pdf", "zh_sample.pdf"):
    p = pymupdf.open(OUT / name)[0]
    p.get_pixmap(dpi=120).save(str(OUT / (name.replace(".pdf", ".png"))))

print(json.dumps(report, ensure_ascii=False, indent=1))
