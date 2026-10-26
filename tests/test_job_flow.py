import io
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from dify_plugin.entities.invoke_message import InvokeMessage
from dify_plugin.entities.tool import ToolInvokeMessage

import tools.parse_document as parse_module
from tools.parse_document import ParseDocumentTool
from tools.render_document import RenderDocumentTool
from tools.translate_shard import TranslateShardTool


DICT = {
    "Revenue": "营收",
    "Operating expenses": "运营费用",
    "Confidential": "机密",
}


class FakeLLM:
    def invoke(self, model_config, prompt_messages, stream=False):
        payload = json.loads(prompt_messages[-1].content)
        items = []
        for item in payload["items"]:
            text = item["text"]
            for src, dst in DICT.items():
                text = text.replace(src, dst)
            items.append({"id": item["id"], "text": text})
        content = json.dumps({"items": items}, ensure_ascii=False)
        return SimpleNamespace(message=SimpleNamespace(content=content))


def build_docx(path: Path, paragraphs: int = 2) -> None:
    doc = Document()
    for i in range(paragraphs):
        doc.add_paragraph(f"Revenue grew by {18 + i} percent.")
    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(path)


def make_tool(cls):
    tool = cls.__new__(cls)
    tool.response_type = ToolInvokeMessage
    tool.runtime = SimpleNamespace(user_id="tester")
    tool.session = SimpleNamespace(model=SimpleNamespace(llm=FakeLLM()))
    return tool


@pytest.fixture()
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("TRANSLATE_BUDDY_HOME", str(tmp_path / "home"))
    return tmp_path


def test_parse_translate_render(workspace, monkeypatch):
    src = workspace / "in.docx"
    build_docx(src)
    monkeypatch.setattr(
        parse_module, "download_to", lambda url, dst, **kw: shutil.copyfile(src, dst)
    )

    parse = make_tool(ParseDocumentTool)
    messages = list(
        parse._invoke(
            {
                "file": SimpleNamespace(url="http://x/in.docx", filename="报告.docx"),
                "target_lang": "zh-Hans",
                "source_lang": "auto",
                "shard_items": 60,
            }
        )
    )
    report = next(m for m in messages if m.type == InvokeMessage.MessageType.JSON).message.json_object
    job_id = report["job_id"]
    assert report["shard_count"] >= 1

    shard = make_tool(TranslateShardTool)
    model = {"provider": "fake", "model": "fake", "mode": "chat"}
    for index in range(report["shard_count"]):
        list(shard._invoke({"job_id": job_id, "shard_index": index, "model": model}))

    # 再跑一次应命中缓存
    again = list(shard._invoke({"job_id": job_id, "shard_index": 0, "model": model}))
    cached = next(m for m in again if m.type == InvokeMessage.MessageType.JSON).message.json_object
    assert cached["cached"] is True

    render = make_tool(RenderDocumentTool)
    out = list(render._invoke({"job_id": job_id}))
    blob = next(m for m in out if m.type == InvokeMessage.MessageType.BLOB)
    assert blob.meta["filename"] == "报告.zh-Hans.docx"
    text = Document(io.BytesIO(blob.message.blob))
    assert "营收" in text.paragraphs[0].text
    assert "机密" in text.sections[0].header.paragraphs[0].text


def test_render_reports_missing_shards(workspace, monkeypatch):
    src = workspace / "in.docx"
    build_docx(src, paragraphs=7)
    monkeypatch.setattr(
        parse_module, "download_to", lambda url, dst, **kw: shutil.copyfile(src, dst)
    )

    parse = make_tool(ParseDocumentTool)
    messages = list(
        parse._invoke(
            {
                "file": SimpleNamespace(url="http://x/in.docx", filename="a.docx"),
                "target_lang": "en",
                "shard_items": 1,
            }
        )
    )
    report = next(m for m in messages if m.type == InvokeMessage.MessageType.JSON).message.json_object
    assert report["shard_count"] >= 2

    render = make_tool(RenderDocumentTool)
    out = list(render._invoke({"job_id": report["job_id"]}))
    result = next(m for m in out if m.type == InvokeMessage.MessageType.JSON).message.json_object
    assert len(result["missing_shards"]) == report["shard_count"]
    assert result["applied"] == 0
