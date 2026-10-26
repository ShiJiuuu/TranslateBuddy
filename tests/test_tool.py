import io
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.shared import RGBColor

from dify_plugin.entities.invoke_message import InvokeMessage
from dify_plugin.entities.tool import ToolInvokeMessage

import tools.translate_document as module
from backend.errors import UnsupportedFormat
from tools.translate_document import TranslateDocumentTool


def build_sample(path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Revenue grew by 18 percent")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p.add_run(" compared with the same period last year, driven by strong demand in the ")
    r2 = p.add_run("enterprise segment")
    r2.italic = True
    p.add_run(".")
    doc.add_paragraph("Operating expenses remained flat.")
    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(path)


class FakeLLM:
    def __init__(self) -> None:
        self.calls = 0

    def invoke(self, model_config, prompt_messages, stream=False):
        self.calls += 1
        payload = json.loads(prompt_messages[-1].content)
        items = []
        for item in payload["items"]:
            text = (
                item["text"]
                .replace("Revenue", "营收")
                .replace("Operating expenses", "运营费用")
                .replace("Confidential", "机密")
            )
            items.append({"id": item["id"], "text": text})
        content = json.dumps({"items": items}, ensure_ascii=False)
        return SimpleNamespace(message=SimpleNamespace(content=content))


def make_tool() -> TranslateDocumentTool:
    tool = TranslateDocumentTool.__new__(TranslateDocumentTool)
    tool.response_type = ToolInvokeMessage
    tool.session = SimpleNamespace(model=SimpleNamespace(llm=FakeLLM()))
    return tool


def test_docx_end_to_end(tmp_path, monkeypatch):
    src = tmp_path / "in.docx"
    build_sample(src)
    monkeypatch.setattr(module, "download_to", lambda url, dst, **kw: shutil.copyfile(src, dst))

    tool = make_tool()
    params = {
        "file": SimpleNamespace(url="http://example.test/in.docx", filename="合同.docx"),
        "target_lang": "zh-Hans",
        "source_lang": "auto",
        "model": {"provider": "fake", "model": "fake", "mode": "chat"},
    }
    messages = list(tool._invoke(params))

    blobs = [m for m in messages if m.type == InvokeMessage.MessageType.BLOB]
    assert len(blobs) == 1
    blob = blobs[0]
    assert blob.meta["filename"] == "合同.zh-Hans.docx"
    assert blob.meta["mime_type"].endswith("wordprocessingml.document")

    result = Document(io.BytesIO(blob.message.blob))
    assert "营收" in result.paragraphs[0].text
    assert "运营费用" in result.paragraphs[1].text
    assert "机密" in result.sections[0].header.paragraphs[0].text

    json_messages = [m for m in messages if m.type == InvokeMessage.MessageType.JSON]
    assert json_messages
    report = json_messages[0].message.json_object
    assert report["segments"] >= 3
    assert report["translated"] >= 3
    assert report["output_filename"] == "合同.zh-Hans.docx"


def test_unsupported_format_is_rejected():
    tool = make_tool()
    params = {
        "file": SimpleNamespace(url="http://example.test/a.pdf", filename="a.pdf"),
        "target_lang": "zh-Hans",
        "model": {"provider": "fake", "model": "fake", "mode": "chat"},
    }
    with pytest.raises(UnsupportedFormat):
        list(tool._invoke(params))
