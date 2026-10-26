import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from backend.xlsx_io import XlsxDocument


def build_workbook(path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Quarterly Report 2024"
    ws["A1"].font = Font(size=16, bold=True)
    for col, text in enumerate(["Metric", "Value", "Note"], start=1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    ws["A4"], ws["B4"], ws["C4"] = "Revenue", 1240, "Enterprise demand"
    ws["A5"], ws["B5"] = "Cost", 880
    ws["A6"], ws["B6"] = "Margin", "=B4-B5"
    ws["B4"].number_format = '#,##0 "USD"'
    wb.save(path)


def test_parse_collects_strings_only(tmp_path):
    src = tmp_path / "in.xlsx"
    build_workbook(src)
    doc = XlsxDocument(src)
    plains = {s.plain for s in doc.segments}
    assert "Quarterly Report 2024" in plains
    assert "Revenue" in plains
    assert "Enterprise demand" in plains
    assert "1,240" not in plains          # 数字不翻
    assert "=B4-B5" not in plains         # 公式不翻


def test_apply_keeps_styles_formula_and_merge(tmp_path):
    src = tmp_path / "in.xlsx"
    build_workbook(src)
    doc = XlsxDocument(src)
    for seg in doc.segments:
        doc.apply(seg.id, "译:" + seg.plain)
    out = doc.save(tmp_path / "out.xlsx")

    wb = load_workbook(out)
    ws = wb["Report"]
    assert ws["A1"].value == "译:Quarterly Report 2024"
    assert ws["A4"].value == "译:Revenue"
    assert ws["B4"].value == 1240
    assert ws["B6"].value == "=B4-B5"
    assert ws["B4"].number_format == '#,##0 "USD"'
    assert ws["A1"].font.bold is True
    assert ws["A3"].fill.fgColor.rgb == "002F5597"
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:C1"]


def test_other_parts_stay_byte_identical(tmp_path):
    src = tmp_path / "in.xlsx"
    build_workbook(src)
    doc = XlsxDocument(src)
    for seg in doc.segments:
        doc.apply(seg.id, "译")
    out = doc.save(tmp_path / "out.xlsx")

    with zipfile.ZipFile(src) as before, zipfile.ZipFile(out) as after:
        assert before.namelist() == after.namelist()
        assert before.read("xl/styles.xml") == after.read("xl/styles.xml")
        assert before.read("xl/workbook.xml") == after.read("xl/workbook.xml")


def test_duplicate_strings_are_detected(tmp_path):
    src = tmp_path / "dup.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Confidential"
    ws["A2"] = "Confidential"
    wb.save(src)

    doc = XlsxDocument(src)
    matches = [s for s in doc.segments if s.plain == "Confidential"]
    assert matches  # 去重发生在翻译层，见 test_translator
