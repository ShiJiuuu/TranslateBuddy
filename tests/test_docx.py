import zipfile
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from backend.docx_io import DocxDocument


def build_sample(path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Revenue grew by 18 percent")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p.add_run(" compared with the same period last year, driven by strong demand in the ")
    r2 = p.add_run("enterprise segment")
    r2.italic = True
    p.add_run(".")

    doc.add_paragraph("Operating expenses remained flat.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "1,240"

    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(path)


def test_parse_covers_body_header_and_table(tmp_path):
    src = tmp_path / "in.docx"
    build_sample(src)

    doc = DocxDocument(src)
    parts = {s.context["part"] for s in doc.segments}
    assert "word/document.xml" in parts
    assert "word/header1.xml" in parts

    plains = [s.plain for s in doc.segments]
    assert "Confidential" in plains
    assert "Metric" in plains
    assert "1,240" not in plains  # 纯数字不翻译

    seg = next(s for s in doc.segments if s.plain.startswith("Revenue grew"))
    assert seg.span_count == 4
    assert seg.text.startswith("[[0]]Revenue grew by 18 percent[[/0]]")


def test_apply_keeps_bold_italic_and_color(tmp_path):
    src = tmp_path / "in.docx"
    build_sample(src)
    doc = DocxDocument(src)
    seg = next(s for s in doc.segments if s.plain.startswith("Revenue grew"))

    translated = (
        "[[0]]营收同比增长 18%[[/0]]"
        "[[1]]，较去年同期，主要来自[[/1]]"
        "[[2]]企业客户[[/2]]"
        "[[3]]。[[/3]]"
    )
    assert doc.apply(seg.id, translated) is True

    out = doc.save(tmp_path / "out.docx")
    result = Document(out)
    runs = [r for r in result.paragraphs[0].runs if r.text]
    assert runs[0].text == "营收同比增长 18%"
    assert runs[0].bold is True
    assert str(runs[0].font.color.rgb) == "C0392B"
    assert runs[2].italic is True
    assert result.paragraphs[0].text == "营收同比增长 18%，较去年同期，主要来自企业客户。"


def test_untagged_translation_falls_back_to_first_run(tmp_path):
    src = tmp_path / "in.docx"
    build_sample(src)
    doc = DocxDocument(src)
    seg = next(s for s in doc.segments if s.plain.startswith("Revenue grew"))

    assert doc.apply(seg.id, "整段译文") is False
    out = doc.save(tmp_path / "out.docx")
    paragraph = Document(out).paragraphs[0]
    assert paragraph.text == "整段译文"
    assert paragraph.runs[0].bold is True  # 格式仍在，只是降级到首片段


def test_other_zip_parts_are_byte_identical(tmp_path):
    src = tmp_path / "in.docx"
    build_sample(src)
    doc = DocxDocument(src)
    seg = next(s for s in doc.segments if s.context["part"] == "word/document.xml")
    doc.apply(seg.id, "译文")
    out = doc.save(tmp_path / "out.docx")

    with zipfile.ZipFile(src) as before, zipfile.ZipFile(out) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name.startswith("word/document"):
                continue
            assert before.read(name) == after.read(name), name


def test_output_still_opens_with_python_docx(tmp_path):
    src = tmp_path / "in.docx"
    build_sample(src)
    doc = DocxDocument(src)
    for seg in list(doc.segments)[:3]:
        doc.apply(seg.id, "译文")
    out = doc.save(tmp_path / "out.docx")
    assert Document(out).paragraphs
